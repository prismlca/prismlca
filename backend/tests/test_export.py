from __future__ import annotations

import io

from openpyxl import load_workbook

from prism_lca.engine.bootstrap import get_bd


def _seed(bd) -> None:
    bd.Database("biosphere").write(
        {("biosphere", "co2"): {"name": "Carbon dioxide", "type": "emission",
                                "categories": ("air",), "unit": "kg"}}
    )
    bd.Database("cfg").write(
        {
            ("cfg", "elec"): {
                "name": "electricity", "unit": "kWh", "reference product": "electricity",
                "exchanges": [
                    {"input": ("cfg", "elec"), "amount": 1, "type": "production"},
                    {"input": ("biosphere", "co2"), "amount": 0.5, "type": "biosphere"},
                ],
            },
            ("cfg", "bottle"): {
                "name": "bottle", "unit": "unit", "reference product": "bottle",
                "exchanges": [
                    {"input": ("cfg", "bottle"), "amount": 1, "type": "production"},
                    {"input": ("cfg", "elec"), "amount": 2, "type": "technosphere"},
                    {"input": ("biosphere", "co2"), "amount": 0.1, "type": "biosphere"},
                ],
            },
        }
    )
    bd.Method(("test", "gwp")).write([(("biosphere", "co2"), 1.0)])


def test_export_lcia_csv(client, token, project):
    _seed(get_bd())
    r = client.post(
        "/api/v1/export/lcia",
        json={"database": "cfg", "code": "bottle", "methods": [["test", "gwp"]], "format": "csv"},
        headers=token,
    )
    assert r.status_code == 200, r.text
    assert "text/csv" in r.headers["content-type"]
    lines = r.text.strip().splitlines()
    assert lines[0] == "method,unit,score"
    assert lines[1].startswith("test · gwp,")
    assert "1.1" in lines[1]


def test_export_contribution_json_and_xlsx(client, token, project):
    _seed(get_bd())
    # JSON returns the raw envelope (lossless)
    rj = client.post(
        "/api/v1/export/contribution",
        json={"database": "cfg", "code": "bottle", "method": ["test", "gwp"], "format": "json"},
        headers=token,
    )
    assert rj.status_code == 200, rj.text
    assert abs(rj.json()["meta"]["total"] - 1.1) < 1e-6

    # XLSX opens as a real workbook with the flattened tree
    rx = client.post(
        "/api/v1/export/contribution",
        json={"database": "cfg", "code": "bottle", "method": ["test", "gwp"], "format": "xlsx"},
        headers=token,
    )
    assert rx.status_code == 200, rx.text
    assert "spreadsheetml" in rx.headers["content-type"]
    wb = load_workbook(io.BytesIO(rx.content))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["depth", "code", "name", "location", "amount",
                                        "unit", "impact", "direct", "share"]
    names = [row[2] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert "bottle" in names
    assert "electricity" in names


def test_export_compare_matrix_xlsx(client, token, project):
    import base64

    r = client.post(
        "/api/v1/export/compare-matrix",
        json={
            "columns": [
                {"key": "a", "family": "EF v3.1", "label": "climate change", "unit": "kg CO2-Eq"},
                {"key": "b", "family": "ReCiPe", "label": "GWP", "unit": "kg CO2-Eq"},
            ],
            "subjects": [
                {"label": "Verarbeitung", "scores": {"a": 0.77, "b": 0.8}},
                {"label": "Alt A", "scores": {"a": 1.2, "b": None}},
            ],
            "filename": "res",
        },
        headers=token,
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["filename"] == "res.xlsx"
    assert base64.b64decode(body["base64"])[:2] == b"PK"


def test_export_compare_matrix_disaggregated(client, token, project):
    """With a group header the sheet gains an identity column and the rows come apart per phase,
    with the row that sums them marked — the shape you can sort, filter and pivot."""
    import base64
    import io

    from openpyxl import load_workbook

    r = client.post(
        "/api/v1/export/compare-matrix",
        json={
            "columns": [
                {"key": "a", "family": "EF v3.1", "label": "climate change", "unit": "kg CO2-Eq"},
            ],
            "group_header": "Life-cycle phase",
            "subjects": [
                {"label": "Bottle", "group": "Production", "scores": {"a": 0.6}},
                {"label": "Bottle", "group": "Use", "scores": {"a": 0.4}},
                {"label": "Bottle", "group": "Total", "scores": {"a": 1.0}, "total": True},
            ],
            "filename": "res",
        },
        headers=token,
    )
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(base64.b64decode(r.json()["base64"]))).active

    assert ws.cell(1, 1).value == "Subject"
    assert ws.cell(1, 2).value == "Life-cycle phase"
    # The indicators start after BOTH identity columns.
    assert ws.cell(2, 3).value.startswith("climate change")
    assert [ws.cell(row, 2).value for row in (3, 4, 5)] == ["Production", "Use", "Total"]
    assert [ws.cell(row, 3).value for row in (3, 4, 5)] == [0.6, 0.4, 1.0]
    # The subject repeats instead of being merged — a merged cell cannot be sorted or pivoted.
    assert [ws.cell(row, 1).value for row in (3, 4, 5)] == ["Bottle"] * 3
    # The summing row stands out.
    assert ws.cell(5, 3).font.bold is True
    assert ws.cell(4, 3).font.bold is not True
    assert ws.freeze_panes == "C3"


def test_export_compare_matrix_without_a_group_is_unchanged(client, token, project):
    import base64
    import io

    from openpyxl import load_workbook

    r = client.post(
        "/api/v1/export/compare-matrix",
        json={
            "columns": [{"key": "a", "family": "EF v3.1", "label": "climate change", "unit": "kg"}],
            "subjects": [{"label": "Bottle", "scores": {"a": 1.0}}],
        },
        headers=token,
    )
    ws = load_workbook(io.BytesIO(base64.b64decode(r.json()["base64"]))).active
    assert ws.cell(1, 1).value == "Subject"
    assert ws.cell(2, 2).value.startswith("climate change")  # indicators still start at B
    assert ws.freeze_panes == "B3"


def test_export_elementary_flows_xlsx(client, token, project):
    import base64

    r = client.post(
        "/api/v1/export/elementary-flows",
        json={
            "scope": "Whole system",
            "method": ["EF v3.1", "climate change", "GWP100"],
            "unit": "kg CO2-Eq",
            "total": 1.23,
            "flows": [
                {"flow": "Carbon dioxide, fossil", "categories": ["air"], "impact": 1.0, "share": 0.81},
                {"flow": "Methane, fossil", "categories": ["air", "urban air"], "impact": 0.2, "share": 0.16},
            ],
        },
        headers=token,
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["filename"] == "prism-elementary-flows.xlsx"
    assert base64.b64decode(body["base64"])[:2] == b"PK"
