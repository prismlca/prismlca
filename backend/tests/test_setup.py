from __future__ import annotations

import pytest

from prism_lca.errors import PrismError

NS = "http://www.EcoInvent.org/EcoSpold02"

# (code, name, compartment, subcompartment) — mirrors the shape of a real MasterData flow list.
FLOWS = [
    ("4505f26e-0000-0000-0000-000000000001", "Flurochloridone", "air", "non-urban air or from high stacks"),
    ("4505f26e-0000-0000-0000-000000000002", "1,2-Dibromoethane", "water", "unspecified"),
    ("4505f26e-0000-0000-0000-000000000003", "Carbon dioxide", "air", "urban air close to ground"),
]


def write_flow_list(path, flows):
    """A minimal but structurally real ``ElementaryExchanges.xml`` (bw2io reads it with lxml)."""
    entries = "\n".join(
        f"""  <elementaryExchange id="{code}" casNumber="000000-00-0">
    <name xml:lang="en">{name}</name>
    <unitName xml:lang="en">kg</unitName>
    <compartment>
      <compartment xml:lang="en">{compartment}</compartment>
      <subcompartment xml:lang="en">{subcompartment}</subcompartment>
    </compartment>
  </elementaryExchange>"""
        for code, name, compartment, subcompartment in flows
    )
    path.write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<validElementaryExchanges xmlns="{NS}" majorRelease="3" minorRelease="0" '
        'majorRevision="0" minorRevision="70158">\n'
        f"{entries}\n</validElementaryExchanges>\n",
        encoding="utf-8",
    )
    return path


def write_lcia_workbook(path, cf_rows, indicator_rows):
    from openpyxl import Workbook

    wb = Workbook()
    cfs = wb.active
    cfs.title = "CFs"
    cfs.append(["Method", "Category", "Indicator", "Name", "Compartment", "Subcompartment", "CF"])
    for row in cf_rows:
        cfs.append(list(row))
    indicators = wb.create_sheet("Indicators")
    indicators.append(["Method", "Category", "Indicator", "Indicator Unit"])
    for row in indicator_rows:
        indicators.append(list(row))
    wb.save(path)
    return path


def test_setup_status_is_reportable(client, token):
    r = client.get("/api/v1/setup/status", headers=token)
    assert r.status_code == 200
    assert r.json()["state"] in {"idle", "running", "done", "error"}
    assert "progress" in r.json()  # 0–1 while measurable, null otherwise — the wizard's bar


def test_ecoinvent_import_rejects_missing_path(client, token):
    r = client.post(
        "/api/v1/setup/ecoinvent",
        json={"ecospold2_dir": "/definitely/not/a/real/path"},
        headers=token,
    )
    assert r.status_code == 400
    body = r.json()
    assert body["error"]["code"] == "path_not_found"
    assert body["error"]["hint"]


def test_ecoinvent_import_rejects_missing_lcia_file(client, token, tmp_path):
    r = client.post(
        "/api/v1/setup/ecoinvent",
        json={"ecospold2_dir": str(tmp_path), "lcia_xlsx": "/nope/methods.xlsx"},
        headers=token,
    )
    assert r.status_code == 400
    assert r.json()["error"]["code"] == "lcia_file_not_found"


def test_lcia_import_rejects_missing_file(client, token):
    r = client.post("/api/v1/setup/lcia", json={"lcia_xlsx": "/nope/methods.xlsx"}, headers=token)
    assert r.status_code == 400
    assert r.json()["error"]["code"] == "lcia_file_not_found"


def test_biosphere_is_built_from_the_release_flow_list(client, token, project, tmp_path):
    """The flows come from the release's own list — never from bw2io's bundled 3.9 one."""
    from prism_lca.engine.bootstrap import get_bd
    from prism_lca.engine.ecoinvent import (
        BIO_FLOW_COUNT_KEY,
        BIO_REVISION_KEY,
        BIOSPHERE_DB,
        _ensure_biosphere,
    )

    note, _flows = _ensure_biosphere(write_flow_list(tmp_path / "ElementaryExchanges.xml", FLOWS))
    assert "3 flows" in note

    bd = get_bd()
    flows = {f["name"]: f for f in bd.Database(BIOSPHERE_DB)}
    assert set(flows) == {"Flurochloridone", "1,2-Dibromoethane", "Carbon dioxide"}
    # "unspecified" subcompartments collapse to a one-element context.
    assert tuple(flows["1,2-Dibromoethane"]["categories"]) == ("water",)
    assert tuple(flows["Flurochloridone"]["categories"]) == (
        "air",
        "non-urban air or from high stacks",
    )
    # Provenance, so a later import can name the flow list instead of guessing.
    assert bd.databases[BIOSPHERE_DB][BIO_REVISION_KEY] == "3.0.0.70158"
    assert bd.databases[BIOSPHERE_DB][BIO_FLOW_COUNT_KEY] == 3


def test_existing_biosphere_is_patched_not_skipped(client, token, project, tmp_path):
    """An older biosphere must gain the release's new flows — skipping it would fail the
    inventory import later with unlinked exchanges and no explanation."""
    from prism_lca.engine.bootstrap import get_bd
    from prism_lca.engine.ecoinvent import BIOSPHERE_DB, _ensure_biosphere

    _ensure_biosphere(write_flow_list(tmp_path / "old.xml", FLOWS[:2]))
    assert len(get_bd().Database(BIOSPHERE_DB)) == 2

    note, _flows = _ensure_biosphere(write_flow_list(tmp_path / "new.xml", FLOWS))
    assert "added 1 missing flows" in note
    assert len(get_bd().Database(BIOSPHERE_DB)) == 3
    assert "Carbon dioxide" in {f["name"] for f in get_bd().Database(BIOSPHERE_DB)}

    # Running it again with the same list changes nothing.
    again, _flows = _ensure_biosphere(write_flow_list(tmp_path / "same.xml", FLOWS))
    assert "already holds all 3 flows" in again
    assert len(get_bd().Database(BIOSPHERE_DB)) == 3


def test_lcia_import_links_dedupes_and_keeps_units(client, token, project, tmp_path):
    from prism_lca.engine.bootstrap import get_bd
    from prism_lca.engine.ecoinvent import _ensure_biosphere, _import_lcia

    _ensure_biosphere(write_flow_list(tmp_path / "flows.xml", FLOWS))
    method = ("Test method", "climate change", "GWP100")
    xlsx = write_lcia_workbook(
        tmp_path / "lcia.xlsx",
        [
            (*method, "Flurochloridone", "air", "non-urban air or from high stacks", 1.5),
            (*method, "1,2-Dibromoethane", "water", "unspecified", 2.0),
            # Same flow reached through the subcompartment fallback — must NOT be added twice,
            # because bw2calc sums duplicate matrix entries.
            (*method, "1,2-Dibromoethane", "water", "river", 9.9),
            (*method, "Carbon dioxide", "air", "urban air close to ground", None),  # no CF
        ],
        [(*method, "kg CO2-Eq")],
    )

    note = _import_lcia(xlsx)
    assert "Wrote 1 LCIA methods" in note
    assert "1 duplicate flow mappings skipped" in note

    bd = get_bd()
    assert method in bd.methods
    assert bd.methods[method]["unit"] == "kg CO2-Eq"
    loaded = dict(bd.Method(method).load())
    assert len(loaded) == 2
    assert sorted(loaded.values()) == [1.5, 2.0]  # the 9.9 duplicate row lost, first wins


def test_lcia_import_refuses_unlinked_cfs_and_names_candidates(client, token, project, tmp_path):
    """Half-linked methods would look plausible and be wrong, so the import must abort — and
    say which flow it could not reach, plus the nearest candidate."""
    from prism_lca.engine.ecoinvent import _ensure_biosphere, _import_lcia

    _ensure_biosphere(write_flow_list(tmp_path / "flows.xml", FLOWS))
    method = ("Test method", "climate change", "GWP100")
    xlsx = write_lcia_workbook(
        tmp_path / "lcia.xlsx",
        [(*method, "Carbon dioxid", "air", "urban air close to ground", 1.0)],
        [(*method, "kg CO2-Eq")],
    )

    with pytest.raises(PrismError) as excinfo:
        _import_lcia(xlsx)
    assert excinfo.value.code == "lcia_unlinked"
    assert "Carbon dioxid" in excinfo.value.hint
    # The hint points at the likely intended flow instead of silently substituting it.
    assert "Carbon dioxide" in excinfo.value.hint

    from prism_lca.engine.bootstrap import get_bd

    assert method not in get_bd().methods


def test_lcia_import_needs_a_biosphere(client, token, project, tmp_path):
    from prism_lca.engine.ecoinvent import _import_lcia

    xlsx = write_lcia_workbook(tmp_path / "lcia.xlsx", [], [])
    with pytest.raises(PrismError) as excinfo:
        _import_lcia(xlsx)
    assert excinfo.value.code == "biosphere_missing"


def _release_dir(tmp_path, flows=FLOWS):
    """A stand-in release: an (empty) datasets folder with MasterData beside it."""
    datasets = tmp_path / "datasets"
    datasets.mkdir(exist_ok=True)
    master = tmp_path / "MasterData"
    master.mkdir(exist_ok=True)
    write_flow_list(master / "ElementaryExchanges.xml", flows)
    return datasets


def test_import_refuses_an_empty_leftover_database(client, token, project, tmp_path):
    """An empty registered database is the signature of a crashed write."""
    from prism_lca.engine import ecoinvent
    from prism_lca.engine.bootstrap import get_bd

    get_bd().Database("ei-leftover").write({})
    ecoinvent.start_import(
        str(_release_dir(tmp_path)), project, "ei-leftover", run_in_thread=False
    )
    st = ecoinvent.status()
    assert st.state is ecoinvent.JobState.error
    assert "is empty" in st.detail


def test_import_accepts_a_database_imported_before_the_completion_flag(
    client, token, project, tmp_path
):
    """A missing completion flag must NOT lock an existing installation out of its own project —
    databases imported before PrismLCA recorded the flag carry their activities perfectly well."""
    from prism_lca.engine import ecoinvent
    from prism_lca.engine.bootstrap import get_bd

    bd = get_bd()
    bd.Database("ei-legacy").write(
        {("ei-legacy", "elec"): {"name": "electricity", "unit": "kWh", "location": "DE",
                                 "reference product": "electricity", "exchanges": []}}
    )
    assert not bd.databases["ei-legacy"].get(ecoinvent.COMPLETE_KEY)

    ecoinvent.start_import(
        str(_release_dir(tmp_path)), project, "ei-legacy", run_in_thread=False
    )
    st = ecoinvent.status()
    assert st.state is ecoinvent.JobState.done, st.detail
    assert "already imported" in st.detail


def _seed_background(bd) -> None:
    """A read-only background + a biosphere + a method + a model that points into both."""
    from prism_lca.engine import guardrails

    bd.Database("biosphere3").write(
        {("biosphere3", "co2"): {"name": "Carbon dioxide", "type": "emission",
                                 "categories": ("air",), "unit": "kg"}}
    )
    bd.Database("bg-test").write(
        {("bg-test", "elec"): {
            "name": "electricity", "unit": "kWh", "reference product": "electricity",
            "location": "DE",
            "exchanges": [
                {"input": ("bg-test", "elec"), "amount": 1, "type": "production"},
                {"input": ("biosphere3", "co2"), "amount": 0.5, "type": "biosphere"},
            ],
        }}
    )
    bd.Database("my-model").write(
        {("my-model", "widget"): {
            "name": "widget", "unit": "unit", "reference product": "widget", "location": "GLO",
            "exchanges": [
                {"input": ("my-model", "widget"), "amount": 1, "type": "production"},
                {"input": ("bg-test", "elec"), "amount": 3, "type": "technosphere"},
                {"input": ("biosphere3", "co2"), "amount": 0.2, "type": "biosphere"},
            ],
        }}
    )
    bd.Method(("test", "gwp")).write([(("biosphere3", "co2"), 1.0)])
    guardrails.set_read_only("bg-test", True)
    guardrails.set_read_only("biosphere3", True)
    guardrails.set_read_only("my-model", False)


def test_background_overview_counts_what_a_reset_would_cost(client, token, project):
    from prism_lca.engine.bootstrap import get_bd

    _seed_background(get_bd())
    body = client.get("/api/v1/setup/background", headers=token).json()

    assert body["biosphere"]["flows"] == 1
    assert [d["name"] for d in body["background"]] == ["bg-test"]
    assert body["methods"] == {"total": 1, "custom": 0}
    model = next(m for m in body["models"] if m["name"] == "my-model")
    # The production exchange stays; the technosphere and biosphere links do not.
    assert model["dangling_exchanges"] == 2


def test_reset_needs_a_code_that_was_issued(client, token, project, monkeypatch):
    from prism_lca.engine import background_reset
    from prism_lca.engine.bootstrap import get_bd

    # The code lives in module state: one issued by an earlier test in the same worker
    # (test_projects' reset-during-import test) would turn the 409 below into a 403.
    monkeypatch.setattr(background_reset, "_reset_code", None)
    _seed_background(get_bd())

    r = client.post("/api/v1/setup/background/reset", json={"code": "1234"}, headers=token)
    assert r.status_code == 409
    assert r.json()["error"]["code"] == "reset_code_missing"

    issued = client.post("/api/v1/setup/background/reset-code", headers=token).json()["code"]
    assert len(issued) == 4 and issued.isdigit()

    wrong = "0000" if issued != "0000" else "1111"
    r = client.post("/api/v1/setup/background/reset", json={"code": wrong}, headers=token)
    assert r.status_code == 403
    assert r.json()["error"]["code"] == "reset_code_invalid"
    # A typo must not cost the code — the same one still works.
    assert "bg-test" in get_bd().databases


def test_reset_removes_background_and_methods_but_keeps_the_model(client, token, project):
    from prism_lca.engine.bootstrap import get_bd

    bd = get_bd()
    _seed_background(bd)
    code = client.post("/api/v1/setup/background/reset-code", headers=token).json()["code"]

    r = client.post("/api/v1/setup/background/reset", json={"code": code}, headers=token)
    assert r.status_code == 200, r.text
    body = r.json()
    assert sorted(body["databases"]) == ["bg-test", "biosphere3"]
    assert body["methods"] == 1

    assert "bg-test" not in bd.databases
    assert "biosphere3" not in bd.databases
    assert ("test", "gwp") not in bd.methods
    # The user's own model survives — with the exchanges that now point nowhere.
    assert "my-model" in bd.databases
    assert len(bd.Database("my-model")) == 1

    # Single use: the same code cannot delete twice.
    again = client.post("/api/v1/setup/background/reset", json={"code": code}, headers=token)
    assert again.status_code == 409
    assert again.json()["error"]["code"] == "reset_code_missing"


def test_lcia_import_rejects_a_foreign_workbook(client, token, project, tmp_path):
    from prism_lca.engine.ecoinvent import _ensure_biosphere, _import_lcia

    _ensure_biosphere(write_flow_list(tmp_path / "flows.xml", FLOWS))

    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active.append(["something", "else"])
    wb.save(tmp_path / "wrong.xlsx")

    with pytest.raises(PrismError) as excinfo:
        _import_lcia(tmp_path / "wrong.xlsx")
    assert excinfo.value.code == "lcia_sheets_missing"
