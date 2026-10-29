from __future__ import annotations

import base64
import io
import uuid

from openpyxl import Workbook, load_workbook


def _fg(client, token) -> str:
    db = "fg-" + uuid.uuid4().hex[:6]
    assert client.post("/api/v1/databases/foreground", json={"name": db}, headers=token).status_code == 200
    return db


def _param(client, token, db, name, **body):
    return client.post("/api/v1/params", json={"database": db, "name": name, **body}, headers=token)


def _scenarios(client, token, db) -> list[dict]:
    return client.get("/api/v1/scenarios", params={"database": db}, headers=token).json()


def _upload(rows: list[list]) -> str:
    """Build an xlsx (first row = header) and return it base64-encoded."""
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode("ascii")


def test_template_lists_parameters(client, token, project):
    db = _fg(client, token)
    assert _param(client, token, db, "mass", amount=2.0).status_code == 200
    assert _param(client, token, db, "eff", amount=10.0).status_code == 200

    r = client.get("/api/v1/scenarios/template", params={"database": db}, headers=token)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(base64.b64decode(r.json()["base64"])), data_only=True)
    ws = wb.active
    grid = list(ws.iter_rows(values_only=True))
    assert grid[0][0] == "parameter"
    names = {row[0] for row in grid[1:]}
    assert {"mass", "eff"} <= names


def test_import_create_and_overwrite(client, token, project):
    db = _fg(client, token)
    _param(client, token, db, "mass", amount=2.0)
    _param(client, token, db, "eff", amount=10.0)

    # First import: creates two scenarios. "High" overrides both; "Low" leaves `eff` blank.
    xlsx = _upload([
        ["parameter", "High", "Low"],
        ["mass", 5, 1],
        ["eff", 20, None],
    ])
    r = client.post("/api/v1/scenarios/import", json={"database": db, "xlsx_base64": xlsx}, headers=token)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["created"] == 2 and body["overwritten"] == 0

    by_name = {s["name"]: s for s in _scenarios(client, token, db)}
    assert by_name["High"]["overrides"] == {"mass": 5.0, "eff": 20.0}
    # Blank cell → no override for `eff` (inherits baseline).
    assert by_name["Low"]["overrides"] == {"mass": 1.0}

    # Second import: same names → overwrite. Blank `mass` keeps the existing 5; 0 for `eff` applies.
    xlsx2 = _upload([
        ["parameter", "High"],
        ["mass", None],
        ["eff", 0],
    ])
    r2 = client.post("/api/v1/scenarios/import", json={"database": db, "xlsx_base64": xlsx2}, headers=token)
    assert r2.status_code == 200, r2.text
    assert r2.json()["created"] == 0 and r2.json()["overwritten"] == 1

    high = {s["name"]: s for s in _scenarios(client, token, db)}["High"]
    assert high["overrides"] == {"mass": 5.0, "eff": 0.0}  # blank kept 5, explicit 0 applied


def test_import_rejects_unknown_parameter(client, token, project):
    db = _fg(client, token)
    _param(client, token, db, "mass", amount=2.0)
    xlsx = _upload([["parameter", "S1"], ["nonexistent", 3]])
    r = client.post("/api/v1/scenarios/import", json={"database": db, "xlsx_base64": xlsx}, headers=token)
    assert r.status_code == 400
    assert r.json()["error"]["code"] == "unknown_parameter"


def test_import_rejects_a_text_cell_that_is_no_valid_formula(client, token, project):
    """A text cell is a formula override; one that names nothing the model has is refused,
    and the message says which cell."""
    db = _fg(client, token)
    _param(client, token, db, "mass", amount=2.0)
    xlsx = _upload([["parameter", "S1"], ["mass", "abc"]])
    r = client.post("/api/v1/scenarios/import", json={"database": db, "xlsx_base64": xlsx}, headers=token)
    assert r.status_code == 400
    err = r.json()["error"]
    assert err["code"] == "formula_error"
    assert "'mass'" in err["message"] and "abc" in err["message"] and "'S1'" in err["message"]


def test_import_takes_a_text_cell_as_a_formula_override(client, token, project):
    db = _fg(client, token)
    _param(client, token, db, "mass", amount=2.0)
    _param(client, token, db, "extra", amount=1.0)
    xlsx = _upload([["parameter", "S1"], ["extra", "3 * mass"]])
    r = client.post("/api/v1/scenarios/import", json={"database": db, "xlsx_base64": xlsx}, headers=token)
    assert r.status_code == 200, r.text
    scen = client.get(f"/api/v1/scenarios?database={db}", headers=token).json()
    assert scen[0]["overrides"] == {"extra": "3 * mass"}
