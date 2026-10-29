"""Phase 0 (concept D12): import BAFU:20XY with the REAL importer into a THROWAWAY Brightway
directory and compare GWP 100 and Ecological Scarcity 2021 totals for every dataset with BAFU's own
"LCIA Results" workbook. Nothing touches the working project.

The throwaway project gets ecoinvent's biosphere and the three method families it needs by
COPYING them out of the live project (read-only access: the live ``BRIGHTWAY2_DIR`` is opened,
records are read, the process then re-points ``BRIGHTWAY2_DIR`` at the scratch directory before
bw2data is used for writing). A one-activity read-only stand-in registered as
``ecoinvent-3.12-cutoff`` makes the importer take the "ecoinvent present" path (S2) without copying
26 533 activities; ``--bundled`` skips the copy and exercises the no-ecoinvent path (S3) instead.

    cd backend
    uv run python scripts/probelauf_bafu.py \\
        --bafu-dir "../databases/BAFU/ecoSpold files" \\
        --results "../databases/BAFU/BAFU-2026 v1_LCIA Results/BAFU-2026 v1 LCIA Results_corrected.xlsx" \\
        --scratch /tmp/prism-bafu-probelauf --out local/flow_mappings/probelauf.json

The live directory is read from ``PRISM_LIVE_BRIGHTWAY2_DIR`` (default
``~/.local/share/Prism LCA/brightway``). Output: distribution of the deviations per method, the
largest outliers with the datasets' names, timings, and the import report — as JSON and as a
short Markdown table on stdout for a results note.
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import statistics
import subprocess
import sys
import time
from pathlib import Path

LIVE_PROJECT = "ecoinvent-3.12-cutoff"
FAMILIES = {
    "GWP": ("IPCC 2021", "climate change", "global warming potential (GWP100)"),
    "UBP": ("Ecological Scarcity 2021", "total", "UBP"),
}
EXCEL_COLUMNS = {"GWP": 4, "UBP": 25}  # 0-based: 'GWP 100 [kg CO2-eq]', 'Total [UBP]'


def dump_live(live_dir: Path, out: Path) -> None:
    """Child process: read biosphere3 + the needed methods from the LIVE project into a JSON file."""
    os.environ["BRIGHTWAY2_DIR"] = str(live_dir)
    import bw2data as bd

    bd.projects.set_current(LIVE_PROJECT)
    flows = [
        {
            "code": n["code"],
            "name": n["name"],
            "categories": list(n["categories"]),
            "unit": n.get("unit", ""),
            "type": n.get("type", "emission"),
            "CAS number": n.get("CAS number", ""),
            "synonyms": list(n.get("synonyms", ())),
        }
        for n in bd.Database("biosphere3")
    ]
    code_of_id = {n.id: n["code"] for n in bd.Database("biosphere3")}
    methods = []
    for m in bd.methods:
        if m[0] not in ("IPCC 2021", "Ecological Scarcity 2021", "EF v3.1"):
            continue
        cfs = []
        for key, amount in bd.Method(m).load():
            code = code_of_id.get(key) if isinstance(key, int) else key[1]
            if code:
                cfs.append([code, amount])
        methods.append({"name": list(m), "unit": bd.methods[m].get("unit", ""), "cfs": cfs})
    out.write_text(json.dumps({"flows": flows, "methods": methods}))
    print(f"[dump] {len(flows)} flows, {len(methods)} methods → {out}", flush=True)


def read_results(xlsx: Path) -> dict[tuple[str, str], dict[str, float]]:
    """``(product, unit) → {GWP, UBP}`` from BAFU's workbook. Product is 'name - location'."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["BAFU_2026 v1"] if "BAFU_2026 v1" in wb.sheetnames else wb.worksheets[1]
    rows = ws.iter_rows(values_only=True)
    next(rows)
    next(rows)
    out = {}
    for row in rows:
        if not row or row[0] is None:
            continue
        product, unit = str(row[0]).strip(), str(row[3] or "").strip()
        vals = {}
        for fam, col in EXCEL_COLUMNS.items():
            try:
                vals[fam] = float(row[col])
            except (TypeError, ValueError):
                pass
        out[(product, unit)] = vals
    return out


#: Written into a scratch directory this script created. Only a directory carrying it is ever wiped —
#: a --scratch that pointed at the live Brightway directory would otherwise have deleted every
#: project in it before the first byte was copied.
SCRATCH_MARKER = ".prism-probelauf-scratch"


def _claim_scratch(scratch: Path, live: Path, *, wipe: bool) -> None:
    scratch = scratch.resolve()
    live = live.resolve()
    if scratch == live or scratch in live.parents or live in scratch.parents:
        sys.exit(f"--scratch {scratch} overlaps the live Brightway directory {live}; pick a temporary path")
    if scratch.exists():
        if any(scratch.iterdir()) and not (scratch / SCRATCH_MARKER).exists():
            sys.exit(f"--scratch {scratch} is not empty and was not made by this script; refusing to touch it")
        if wipe:
            shutil.rmtree(scratch)
    scratch.mkdir(parents=True, exist_ok=True)
    (scratch / SCRATCH_MARKER).write_text("made by probelauf_bafu.py; safe to delete\n", encoding="utf-8")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--bafu-dir", type=Path, required=True)
    ap.add_argument("--results", type=Path, required=True)
    ap.add_argument("--scratch", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument(
        "--bundled", action="store_true", help="S3: no ecoinvent stand-in, bundled 3.9 material"
    )
    ap.add_argument("--keep", action="store_true", help="keep the scratch directory")
    ap.add_argument(
        "--reuse",
        action="store_true",
        help="scratch already holds the imported project (skip seed + import, only compare)",
    )
    args = ap.parse_args()

    live = Path(
        os.environ.get(
            "PRISM_LIVE_BRIGHTWAY2_DIR", Path.home() / ".local/share/Prism LCA/brightway"
        )
    )
    scratch: Path = args.scratch
    _claim_scratch(scratch, live, wipe=not args.keep and not args.reuse)
    t0 = time.time()
    timings: dict[str, float] = {}

    live_dump = scratch / "live.json"
    if not args.bundled and not args.reuse:
        subprocess.run(
            [sys.executable, __file__, "--dump-live", str(live), str(live_dump)], check=True
        )
        timings["dump_live_s"] = round(time.time() - t0, 1)

    os.environ["BRIGHTWAY2_DIR"] = str(scratch)
    os.environ.setdefault("PRISM_DATA_DIR", str(scratch / "prism"))
    from prism_lca.engine import ecoinvent
    from prism_lca.engine.backgrounds import common
    from prism_lca.engine.bootstrap import get_bd
    from prism_lca.engine.guardrails import set_read_only

    bd = get_bd()
    bd.projects.set_current("probelauf")
    if not args.bundled and not args.reuse:
        data = json.loads(live_dump.read_text())
        t = time.time()
        bd.Database("biosphere3").write(
            {
                ("biosphere3", f["code"]): {**f, "categories": tuple(f["categories"])}
                for f in data["flows"]
            }
        )
        meta = bd.databases["biosphere3"]
        meta[ecoinvent.BIO_SOURCE_KEY] = "copied from the live project (probelauf)"
        meta[ecoinvent.BIO_REVISION_KEY] = "3.12"
        bd.databases.flush()
        for m in data["methods"]:
            method = bd.Method(tuple(m["name"]))
            method.register(unit=m["unit"], description="copied from the live project (probelauf)")
            method.write([(("biosphere3", code), amount) for code, amount in m["cfs"]])
        # The stand-in: makes `ecoinvent_present()` true so BAFU maps onto the 3.12 list.
        bd.Database(LIVE_PROJECT).write(
            {
                (LIVE_PROJECT, "standin"): {
                    "name": "ecoinvent stand-in (probelauf)",
                    "unit": "unit",
                    "reference product": "stand-in",
                    "location": "GLO",
                    "exchanges": [
                        {"input": (LIVE_PROJECT, "standin"), "amount": 1, "type": "production"}
                    ],
                }
            }
        )
        bd.databases[LIVE_PROJECT][common.COMPLETE_KEY] = True
        bd.databases.flush()
        set_read_only(LIVE_PROJECT, True)
        set_read_only("biosphere3", True)
        timings["seed_s"] = round(time.time() - t, 1)
        print(
            f"[seed] biosphere {len(data['flows'])} flows, {len(data['methods'])} methods ({timings['seed_s']}s)",
            flush=True,
        )

    t = time.time()
    if args.reuse:
        database = next(n for n in bd.databases if n.startswith("bafu-"))
        print(f"[import] reusing {database} ({len(bd.Database(database))} datasets)", flush=True)
    else:
        ecoinvent.start_bafu_import(str(args.bafu_dir), "probelauf", run_in_thread=False)
        st = ecoinvent.status()
        timings["import_s"] = round(time.time() - t, 1)
        print(
            f"[import] {st.state}: {st.step} — {st.detail[:300]} ({timings['import_s']}s)",
            flush=True,
        )
        if st.state is not ecoinvent.JobState.done:
            raise SystemExit(1)
        database = st.database
    info = common.registry(database)

    # ---- LCIA for every dataset ---------------------------------------------------------------
    import bw2calc as bc

    from prism_lca.engine.backgrounds.flow_mapping import UNIT_ALIASES

    t = time.time()
    db = bd.Database(database)
    acts = list(db)
    methods = list(bd.methods)
    gwp = [
        m
        for m in methods
        if m[0] == "IPCC 2021"
        and m[1] in ("climate change: total (excl. biogenic CO2)", "climate change")
        and m[2] == "global warming potential (GWP100)"
    ]
    ubp = [m for m in methods if m[0] == "Ecological Scarcity 2021" and m[1] == "total"]
    if not gwp or not ubp:
        raise SystemExit(f"methods not found: IPCC GWP100 {gwp}, UBP total {ubp}")
    method_names = {"GWP": gwp[0], "UBP": ubp[0]}
    # Every Ecological Scarcity sub-category that has a column in BAFU's workbook — the
    # diagnostic: WHICH category explains a deviation of the total.
    header = next(
        __import__("openpyxl")
        .load_workbook(args.results, read_only=True)["BAFU_2026 v1"]
        .iter_rows(min_row=2, max_row=2, values_only=True)
    )
    columns = {str(h).lower().replace("[ubp]", "").strip(): i for i, h in enumerate(header) if h}
    aliases = {
        "air pollutants and pm": "main air pollutants and pm",
        "ozone depletion": "ozone layer depletion",
        "global warming potential (gwp100)": "global warming",
    }
    for m in methods:
        if m[0] != "Ecological Scarcity 2021" or m[1] == "total":
            continue
        key = aliases.get(m[2].lower(), m[2].lower())
        if key in columns:
            method_names[f"UBP:{m[2]}"] = m
            EXCEL_COLUMNS[f"UBP:{m[2]}"] = columns[key]
    for fam, name in method_names.items():
        print(f"[method] {fam}: {name}", flush=True)
    lca = bc.LCA({acts[0]: 1}, method=method_names["GWP"])
    lca.lci(factorize=True)
    lca.lcia()
    scores: dict[str, dict[str, float]] = {fam: {} for fam in method_names}
    cfs = {}
    for fam, mname in method_names.items():
        lca.switch_method(mname)
        cfs[fam] = lca.characterization_matrix.copy()
    for i, act in enumerate(acts):
        lca.redo_lci({act.id: 1})
        for fam in method_names:
            scores[fam][act["code"]] = float((cfs[fam] * lca.inventory).sum())
        if (i + 1) % 2000 == 0:
            print(f"[lcia] {i + 1}/{len(acts)} ({time.time() - t:.0f}s)", flush=True)
    timings["lcia_s"] = round(time.time() - t, 1)

    # ---- compare --------------------------------------------------------------------------------
    # BAFU's workbook reports per SI-ish unit, not per the dataset's ecoSpold unit: kWh datasets
    # per MJ, hectares per m², km per m, metre-years per km·a. Ours are per dataset unit, so the
    # score is converted to the workbook's unit before comparing (ours per kWh ÷ 3.6 = per MJ).
    to_workbook = {
        ("kWh", "MJ"): 1 / 3.6,
        ("ha", "m2"): 1e-4,
        ("km", "m"): 1e-3,
        ("my", "km*a"): 1000.0,
    }
    reference = read_results(args.results)
    short_units = {v: k for k, v in UNIT_ALIASES.items()}
    comparison: dict[str, list[dict]] = {fam: [] for fam in method_names}
    unmatched = 0
    per_dataset = []
    for act in acts:
        product = f"{act['name']} - {act.get('location', '')}".rstrip()
        unit_short = short_units.get(act["unit"], act["unit"])
        ref = reference.get((product, unit_short))
        their_unit = unit_short
        if ref is None:
            found = [(u, v) for (p, u), v in reference.items() if p == product]
            if found:
                their_unit, ref = found[0]
        if ref is None:
            unmatched += 1
            continue
        factor = to_workbook.get((unit_short, their_unit), 1.0)
        row = {
            "code": act["code"],
            "name": product,
            "unit": unit_short,
            "workbook_unit": their_unit,
        }
        for fam in method_names:
            if fam not in ref:
                continue
            ours, theirs = scores[fam][act["code"]] * factor, ref[fam]
            rel = (
                (ours - theirs) / abs(theirs)
                if theirs
                else (0.0 if abs(ours) < 1e-12 else float("inf"))
            )
            comparison[fam].append(
                {"code": act["code"], "name": product, "ours": ours, "theirs": theirs, "rel": rel}
            )
            row[f"{fam}.ours"], row[f"{fam}.theirs"] = ours, theirs
        per_dataset.append(row)

    summary = {}
    lines = [
        "| Methode | Datensätze | median |Abw.| | ≤ 1 % | ≤ 2 % | ≤ 5 % | ≤ 10 % | > 10 % | Σ ours/Σ theirs |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for fam, rows in comparison.items():
        rels = [abs(r["rel"]) for r in rows if r["rel"] != float("inf")]
        n = len(rows)
        within = {k: sum(1 for r in rels if r <= k) for k in (0.01, 0.02, 0.05, 0.10)}
        ours_sum = sum(r["ours"] for r in rows)
        theirs_sum = sum(r["theirs"] for r in rows)
        ratio = ours_sum / theirs_sum if theirs_sum else float("nan")
        summary[fam] = {
            "datasets": n,
            "median_abs_rel": statistics.median(rels) if rels else None,
            "within_1pct": within[0.01] / n if n else None,
            "within_2pct": within[0.02] / n if n else None,
            "within_5pct": within[0.05] / n if n else None,
            "within_10pct": within[0.10] / n if n else None,
            "sum_ours": ours_sum,
            "sum_theirs": theirs_sum,
            "sum_ratio": ratio,
            "worst": sorted(
                rows, key=lambda r: -abs(r["rel"]) if r["rel"] != float("inf") else -1e9
            )[:50],
        }
        lines.append(
            f"| {fam} | {n} | {statistics.median(rels) * 100:.2f} % | {within[0.01] / n:.1%} | "
            f"{within[0.02] / n:.1%} | {within[0.05] / n:.1%} | {within[0.10] / n:.1%} | "
            f"{(n - within[0.10]) / n:.1%} | {ratio:.3f} |"
            if n
            else f"| {fam} | 0 | | | | | | | |"
        )
    timings["total_s"] = round(time.time() - t0, 1)
    result = {
        "database": database,
        "registry": info.as_dict() if info else None,
        "datasets": len(acts),
        "unmatched_in_workbook": unmatched,
        "methods": {k: list(v) for k, v in method_names.items()},
        "summary": summary,
        "timings": timings,
        "bundled": args.bundled,
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, indent=1, ensure_ascii=False))
    import csv

    with args.out.with_suffix(".csv").open("w", newline="", encoding="utf-8") as fh:
        fields = sorted({k for row in per_dataset for k in row}, key=lambda k: (k.count("."), k))
        w = csv.DictWriter(fh, fieldnames=fields)
        w.writeheader()
        w.writerows(per_dataset)
    print("\n".join(lines))
    print(
        f"[out] {args.out} (+ .csv per dataset) — {len(acts)} datasets, {unmatched} not found in "
        f"the workbook, timings {timings}"
    )
    for fam in ("GWP", "UBP"):
        print(f"\n[worst {fam}]")
        for r in summary[fam]["worst"][:12]:
            print(
                f"  {r['rel'] * 100:+9.1f} %  ours {r['ours']:.4g}  theirs {r['theirs']:.4g}  {r['name']}"
            )


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--dump-live":
        dump_live(Path(sys.argv[2]), Path(sys.argv[3]))
    else:
        main()
