"""Scenarios: named sets of parameter-value overrides (ADR-013, scenario scope), **model-scoped**.

A scenario belongs to a foreground model (database) and overrides that model's parameters; computing
a scenario applies the overrides on top of the baseline, recalculates parameterised exchanges, runs
LCIA, then restores the baseline.

An override is a number or a **formula** (a string, evaluated in the model's parameter namespace).
Either one replaces the parameter's own formula for the length of the window — otherwise the
recalculation would put the formula's result straight back and the override would do nothing — and the
restore puts the formula back. Stored as ``{name: float | str}``; a scenario of
numbers only is stored exactly as before. (A Superstructure/datapackage optimisation can
replace apply/restore later; the result is identical.)

Scenarios are addressed by a stable ``id`` (globally unique) so the ``name`` is a freely renamable
label; the ``name`` is only unique *within a model*.
"""
from __future__ import annotations

import io
from contextlib import contextmanager

from openpyxl import Workbook, load_workbook

from prism_lca.engine import lca, param_window, params, scenario_lookup
from prism_lca.engine._store import JsonFileStore
from prism_lca.engine.model import new_uid
from prism_lca.errors import PrismError, scenario_calc_error


class ScenarioStore(JsonFileStore):
    FILENAME = "scenarios.json"

    def _load(self) -> dict[str, dict]:
        raw = super()._load()
        migrated, changed = self._migrate(raw)
        if changed:
            self._write(migrated)
        return migrated

    def _migrate(self, raw: dict) -> tuple[dict[str, dict], bool]:
        """Upgrade the legacy name-keyed format ({name: {description, overrides}}) to id-keyed
        ({id: {id, name, description, overrides, database?}}). Idempotent. The per-model
        ``database`` tag is added later by params.migrate_project_to_models()."""
        if all(isinstance(v, dict) and v.get("id") == k for k, v in raw.items()):
            return raw, False
        out: dict[str, dict] = {}
        for key, v in raw.items():
            gid = v.get("id") if isinstance(v, dict) and v.get("id") else new_uid()
            out[gid] = {
                "id": gid,
                "name": v.get("name", key) if isinstance(v, dict) else key,
                "description": (v.get("description", "") if isinstance(v, dict) else ""),
                "overrides": (v.get("overrides", {}) if isinstance(v, dict) else {}),
                "database": (v.get("database") if isinstance(v, dict) else None),
            }
        return out, True

    def _write(self, data: dict) -> None:
        self._save(data)

    def list_all(self, database: str) -> list[dict]:
        return [e for e in self._load().values() if e.get("database") == database]

    def get(self, scenario_id: str, database: str | None = None) -> dict:
        """Fetch a scenario by its (globally unique) id — model-agnostic. ``database`` only shapes the refusal:
        the model whose scenarios it lists (engine/scenario_lookup.py)."""
        data = self._load()
        if scenario_id not in data:
            raise scenario_lookup.not_found(data, [scenario_id], database)
        return data[scenario_id]

    def require(self, database: str, scenario_ids: list[str]) -> list[dict]:
        """Every scenario of ``scenario_ids``, checked before anything is calculated — one refusal for all
        the ids that are not there (engine/scenario_lookup.py)."""
        return scenario_lookup.require(self._load(), scenario_ids, database)

    # No mutator here takes ``lca._LCA_LOCK``. They run inside the command bus's lock, and the parameter
    # routes take the LCA lock BEFORE the bus — a store that took it inside the bus deadlocked against
    # them. The order is LCA lock → bus lock → store lock, and it is
    # the scenario ROUTES (api/v1/scenarios.py) that take the LCA lock, around validation and the bus
    # call, so the validation inside these mutators reads the baseline, not an applied scenario.

    def _name_clash(self, data: dict, name: str, database: str, ignore_id: str | None = None) -> bool:
        return any(e["id"] != ignore_id and e.get("database") == database
                   and e["name"].casefold() == name.casefold() for e in data.values())

    def create(
        self,
        database: str,
        name: str,
        description: str = "",
        overrides: dict | None = None,
        scenario_id: str | None = None,
    ) -> dict:
        """``scenario_id``: keep this id if it is free (a bundle import, whose reports name their
        variants by it); a taken or missing id gets a fresh one."""
        with self._lock:
            name = (name or "").strip()
            if not name:
                raise PrismError("Scenario name must not be empty.", code="invalid_name",
                                 hint="Provide a non-empty scenario name.")
            data = self._load()
            if self._name_clash(data, name, database):
                raise PrismError(f"A scenario named '{name}' already exists.", code="scenario_exists",
                                 status_code=409, hint="Choose another name.")
            sid = str(scenario_id) if scenario_id and str(scenario_id) not in data else new_uid()
            entry = {"id": sid, "name": name, "description": description or "",
                     "overrides": check_overrides(database, overrides or {}, scenario=name), "database": database}
            data[entry["id"]] = entry
            self._write(data)
            return entry

    def update(self, scenario_id: str, name: str, description: str, overrides: dict) -> dict:
        with self._lock:
            data = self._load()
            if scenario_id not in data:
                raise scenario_lookup.not_found(data, [scenario_id], None)
            database = data[scenario_id].get("database")
            name = (name or "").strip()
            if not name:
                raise PrismError("Scenario name must not be empty.", code="invalid_name")
            if self._name_clash(data, name, database, ignore_id=scenario_id):
                raise PrismError(f"A scenario named '{name}' already exists.", code="scenario_exists", status_code=409)
            data[scenario_id] = {"id": scenario_id, "name": name, "description": description or "",
                                 "overrides": check_overrides(database, overrides or {}, scenario=name),
                                 "database": database}
            self._write(data)
            return data[scenario_id]

    def duplicate(self, scenario_id: str, new_id: str | None = None) -> dict:
        """``new_id``: keep this id when it is free — a redo recreating the copy it undid, so a later step
        naming the copy by id still finds it."""
        with self._lock:
            src = self.get(scenario_id)
            database = src.get("database")
            data = self._load()
            base = f"{src['name']} (copy)"
            name, i = base, 2
            while self._name_clash(data, name, database):
                name = f"{base} {i}"
                i += 1
            sid = new_id if new_id and new_id not in data else new_uid()
            entry = {"id": sid, "name": name, "description": src.get("description", ""),
                     "overrides": dict(src.get("overrides", {})), "database": database}
            data[entry["id"]] = entry
            self._write(data)
            return entry

    def import_columns(self, database: str, columns: list[dict], ids: dict[str, str] | None = None) -> dict:
        """Create or merge a parsed workbook's columns in ONE write — all or nothing.

        Every merged scenario is checked (``check_overrides``) before anything is stored: a column that
        passes on its own can fail once merged into the existing overrides of its scenario (``b = "4*a"``
        stored, the workbook sets ``a``), and checking while writing left the earlier columns stored by
        a command that then failed — no undo line, no audit entry. Two
        columns for one name merge in order, as they always did.

        Returns ``{"created_ids", "updated_before", "created", "overwritten", "scenarios"}`` — the undo's
        inputs, and the model's scenario list as written, taken from the transaction itself: a read after
        the commit that failed would fail a command whose write already stands, and a failed command is
        not on the undo stack.
        """
        with self._lock:
            data = self._load()
            by_name = {
                e["name"].casefold(): e["id"] for e in data.values() if e.get("database") == database
            }
            staged: dict[str, dict] = {}
            created_ids: list[str] = []
            updated_before: list[dict] = []
            created = overwritten = 0
            for col in columns:
                name = (col.get("name") or "").strip()
                if not name:
                    continue
                overrides = {str(k): as_override(v) for k, v in (col.get("overrides") or {}).items()}
                key = name.casefold()
                sid = by_name.get(key)
                if sid is not None:
                    cur = staged.get(sid) or data[sid]
                    if sid not in created_ids and all(b["id"] != sid for b in updated_before):
                        updated_before.append(dict(data[sid]))
                    staged[sid] = {**cur, "overrides": {**(cur.get("overrides") or {}), **overrides}}
                    overwritten += 1
                else:
                    # A redo reuses the id the first run gave this name (``ids``), when it is free in
                    # the store and among this run's own new ones — later steps name the scenario by it.
                    wanted = (ids or {}).get(key)
                    sid = wanted if wanted and wanted not in data and wanted not in staged else new_uid()
                    by_name[key] = sid
                    created_ids.append(sid)
                    staged[sid] = {"id": sid, "name": name, "description": "", "overrides": overrides,
                                   "database": database}
                    created += 1
            # The whole staged state checked first; the store is written only when all of it holds.
            for entry in staged.values():
                entry["overrides"] = check_overrides(database, entry["overrides"], scenario=entry["name"])
            data.update(staged)
            self._write(data)
            return {"created_ids": created_ids, "updated_before": updated_before,
                    "created": created, "overwritten": overwritten,
                    "scenarios": [e for e in data.values() if e.get("database") == database]}

    def undo_import(self, created_ids: list[str], updated_before: list[dict]) -> None:
        """Take an import back in ONE write: its created scenarios out, the overwritten ones as they were.

        Two steps (delete, then put each) could fail in the middle and leave the store half reverted,
        with the command already off the undo stack."""
        with self._lock:
            data = self._load()
            for sid in created_ids:
                data.pop(sid, None)
            for before in updated_before:
                data[before["id"]] = before
            self._write(data)

    def put(self, scenario_id: str, entry: dict) -> None:
        with self._lock:
            data = self._load()
            data[scenario_id] = entry
            self._write(data)

    def delete(self, scenario_id: str) -> None:
        with self._lock:
            data = self._load()
            if scenario_id in data:
                data.pop(scenario_id)
                self._write(data)

    def set_value(self, scenario_id: str, param: str, value: float | str) -> None:
        with self._lock:
            data = self._load()
            if scenario_id not in data:
                raise scenario_lookup.not_found(data, [scenario_id], None)
            entry = data[scenario_id]
            # Bound by the parameter's own rule, like the parameter itself — a scenario is a set of its values.
            # With the scenario's other overrides in place: a formula override is evaluated against them.
            checked = check_overrides(
                entry.get("database") or "",
                {**(entry.get("overrides") or {}), param: value},
                scenario=entry.get("name", ""),
            )
            entry["overrides"][param] = checked[param]
            self._write(data)

    def remove_value(self, scenario_id: str, param: str) -> None:
        with self._lock:
            data = self._load()
            if scenario_id in data:
                data[scenario_id]["overrides"].pop(param, None)
                self._write(data)

    # --------------------------------------------------------------------- #
    # Model lifecycle
    # --------------------------------------------------------------------- #
    def copy(self, src: str, dst: str) -> None:
        """Copy every scenario of model ``src`` into ``dst`` with fresh ids (model duplicate)."""
        with self._lock:
            data = self._load()
            for e in [e for e in data.values() if e.get("database") == src]:
                nid = new_uid()
                data[nid] = {"id": nid, "name": e["name"], "description": e.get("description", ""),
                             "overrides": dict(e.get("overrides", {})), "database": dst}
            self._write(data)

    def rename_database(self, old: str, new: str) -> None:
        with self._lock:
            data = self._load()
            changed = False
            for e in data.values():
                if e.get("database") == old:
                    e["database"] = new
                    changed = True
            if changed:
                self._write(data)

    def remove_database(self, database: str) -> None:
        with self._lock:
            data = self._load()
            keys = [k for k, e in data.items() if e.get("database") == database]
            if keys:
                for k in keys:
                    data.pop(k, None)
                self._write(data)

    # --------------------------------------------------------------------- #
    # Legacy (project-global) migration helpers
    # --------------------------------------------------------------------- #
    def _legacy_untagged(self) -> list[dict]:
        return [e for e in self._load().values() if not e.get("database")]

    def _drop_legacy_untagged(self) -> None:
        with self._lock:
            data = self._load()
            keys = [k for k, e in data.items() if not e.get("database")]
            if keys:
                for k in keys:
                    data.pop(k, None)
                self._write(data)


def as_override(value: object) -> float | str:
    """One override value in its stored form: a number as a float, a formula as its stripped text.

    A string that reads as a number is a number — the API has always accepted ``"3"`` for 3, and a
    formula ``"3"`` would mean the same thing by a longer road. Raises ``ValueError`` for anything else
    (an empty string, a list …)."""
    if isinstance(value, str):
        text = value.strip()
        if text.startswith("="):  # typed the spreadsheet way
            text = text[1:].strip()
        if not text:
            raise ValueError("an empty override")
        try:
            return float(text)
        except ValueError:
            return text
    if isinstance(value, (int, float)):
        return float(value)
    raise ValueError(f"not a number or a formula: {value!r}")


def check_overrides(database: str, overrides: dict, *, scenario: str = "") -> dict[str, float | str]:
    """Every override checked against its parameter's own rule, and returned in stored form.

    A scenario is a set of parameter values, so it is bound by what those parameters allow: a recycling rate
    a modeller cannot type into the parameter must not arrive through a scenario either — nor through a
    workbook, which is the path where nobody is watching.

    A formula override is evaluated here, against the model as it stands and the scenario's other
    overrides, and refused with the reason when it cannot be (an unknown name, a cycle, a syntax error);
    its *result* is then held to the parameter's rule like a typed number.

    An override naming a parameter this model does not have is left alone here; the import checks that
    separately and says so better.

    Takes no lock itself: it runs inside commands, under the bus lock (see ``ScenarioStore``). The
    caller that wants the baseline rather than an applied scenario holds ``lca._LCA_LOCK`` around it —
    the scenario routes do.
    """
    from prism_lca.engine import param_rules, params

    rules = {p["name"]: (p.get("rule") or {}) for p in params.list_params(database)} if database else {}
    out: dict[str, float | str] = {}
    where = f"in scenario '{scenario}'" if scenario else "in this scenario"
    for name, value in (overrides or {}).items():
        try:
            v = as_override(value)
        except ValueError as exc:
            raise PrismError(
                f"The override for '{name}' {where} is neither a number nor a formula: {value!r}.",
                code="invalid_override",
                hint="Give a number, or a formula over the model's parameters as text (e.g. '2 * mass_kg').",
            ) from exc
        if not isinstance(v, str):
            param_rules.check_value(rules.get(name) or {}, v, name=name, where=where)
        out[name] = v
    formulas = [name for name, v in out.items() if isinstance(v, str)]
    if formulas and not database:
        raise PrismError(
            f"A formula override {where} needs the model it belongs to.",
            code="invalid_override",
            hint="Create the scenario on a model first.",
        )
    if out and database:
        # Numbers too: a number can break a formula of the model (c = 1/(2-a) with a = 2), and a set
        # the model cannot be evaluated under is refused rather than stored to fail at compute time.
        # Not when the model cannot be evaluated even WITHOUT the overrides — that is not this
        # scenario's doing, and a numeric scenario on it was always accepted. Known limit: a broken
        # baseline then also hides an error this scenario adds on top.
        namespace = params.parameter_namespace(database)
        try:
            values = params.evaluate_overrides(database, out, namespace)
        except PrismError as exc:
            if not formulas and _evaluation_error(database, {}, namespace) is not None:
                return out
            exc.message = f"{exc.message.rstrip('.')} ({where})."
            raise
        for name in formulas:
            param_rules.check_value(rules.get(name) or {}, values[name], name=name, where=where)
    return out


class Plan:
    """What applying one scenario writes into its model — computed once, used by every window.

    ``numbers``: the numeric overrides. ``formulas``: the formula every parameter of the model carries
    inside the window (the baseline's, with this scenario's changes — a full map, so a window that moves
    from one scenario to the next puts back what the previous one took out). ``replaced``: the parameters
    whose formula changes, ``{name: {"baseline": old, "wrote": new}}``, for the crash marker.
    ``formula_overrides``: the overrides that are formulas, whose results are checked after the write.
    """

    def __init__(self, overrides: dict[str, float | str], base_formulas: dict[str, str | None]) -> None:
        self.numbers = {k: v for k, v in overrides.items() if not isinstance(v, str)}
        self.formula_overrides = {k: v for k, v in overrides.items() if isinstance(v, str)}
        self.replaced: dict[str, dict] = {}
        for name, value in overrides.items():
            if name not in base_formulas:
                continue  # a parameter this model does not have — the write skips it too
            wrote = value if isinstance(value, str) else None
            if wrote != base_formulas.get(name):
                self.replaced[name] = {"baseline": base_formulas.get(name), "wrote": wrote}
        self.formulas = {**base_formulas, **{k: v["wrote"] for k, v in self.replaced.items()}}

    def replaces_formula(self) -> dict[str, str]:
        """Overridden parameter → the formula of its own that the override replaces."""
        return {k: v["baseline"] for k, v in sorted(self.replaced.items()) if v["baseline"]}


def plan(scenario: dict, base_formulas: dict[str, str | None]) -> Plan:
    """The :class:`Plan` of a stored scenario; a value that is neither a number nor a formula (a store
    edited by hand) is the scenario's calculation error, as it always was."""
    try:
        overrides = {str(k): as_override(v) for k, v in (scenario.get("overrides") or {}).items()}
    except (TypeError, ValueError) as exc:
        raise scenario_calc_error(scenario.get("name") or "?", exc) from exc
    return Plan(overrides, base_formulas)


def check_applied(database: str, applied: Plan, scenario_name: str) -> None:
    """Hold what a formula override computed, now that it is in the model, to its parameter's rule.

    Checked when the scenario was stored too — against the model as it stood then. A baseline edited
    since can move a formula override's result across the rule, and a number the parameter could never
    be given must not flow into a result either. Raises; the window's restore runs in its ``finally``."""
    if not applied.formula_overrides:
        return
    from prism_lca.engine import param_rules

    rules = {p["name"]: (p.get("rule") or {}) for p in params.list_params(database)}
    live = params.snapshot_all(database)
    for name in sorted(applied.formula_overrides):
        if rules.get(name) and name in live:
            param_rules.check_value(rules[name], live[name], name=name, where=f"in scenario '{scenario_name}'")


def _baseline_values(database: str) -> dict[str, float]:
    """The model's parameters as they stand — what a scenario is a variation *of*."""
    from prism_lca.engine import params

    out: dict[str, float] = {}
    for p in params.list_params(database) if database else []:
        if isinstance(p.get("amount"), (int, float)):
            out[str(p["name"])] = float(p["amount"])
    return out


def _same_expression(a: str, b: str) -> bool:
    import ast

    try:
        return ast.dump(ast.parse(a, mode="eval")) == ast.dump(ast.parse(b, mode="eval"))
    except SyntaxError:
        return a.strip() == b.strip()


def ineffective_overrides(database: str, overrides: dict, *, only: set[str] | None = None) -> list[str]:
    """Overrides that set a parameter to the value it already has — a variant that varies nothing.

    Two ways to get here, and the second is why this is a *report* on a scenario rather than only a
    refusal on the way in. Someone types the current value; or the scenario was meaningful and the
    baseline later moved to meet it. An apple-juice model held the second kind: "Longer DC
    distance" sets `dist_dc_store` to 500 km, and the parameter says 500 km. Every comparison returned
    Δ 0, the report drew three bars of the same height, and nothing anywhere said that one of the
    three was the baseline under another name.

    ``only`` judges just those names — within the whole set, which is what decides for a formula
    parameter — and skips the evaluation for every other one (``/value`` names one override).
    """
    with lca._LCA_LOCK:  # the baseline, not a scenario another thread holds applied
        return _Judge(database).ineffective(overrides or {}, only)


class _Judge:
    """One read of a model's parameter state, for judging one or many override sets against it.

    The scenario list judged every scenario from scratch — a namespace read and a full evaluation of
    the model each, and a second one for the evaluation error: 30 scenarios on a 501-parameter model
    took 2 s under the global calculation lock. Here the baseline
    amounts, formulas and namespace are read once, each override set is evaluated once (and that one
    evaluation answers both "does it evaluate?" and the with-all half of "does this override change
    anything?"), and the baseline itself is evaluated only if some set fails.
    """

    _UNSET = object()

    def __init__(self, database: str) -> None:
        self.database = database
        self.baseline = _baseline_values(database)
        self.formulas = params.snapshot_formulas(database) if database else {}
        self._namespace = None
        self._baseline_error = self._UNSET

    @property
    def namespace(self):
        if self._namespace is None:
            self._namespace = params.parameter_namespace(self.database)
        return self._namespace

    def evaluate(self, overrides: dict) -> tuple[dict[str, float] | None, str | None]:
        """(values, None) or (None, why)."""
        try:
            normalised = {str(k): as_override(v) for k, v in (overrides or {}).items()}
            return params.evaluate_overrides(self.database, normalised, self.namespace), None
        except ValueError as exc:
            return None, str(exc)
        except PrismError as exc:
            return None, exc.message

    def baseline_broken(self) -> bool:
        if self._baseline_error is self._UNSET:
            self._baseline_error = self.evaluate({})[1]
        return self._baseline_error is not None

    def ineffective(
        self, overrides: dict, only: set[str] | None = None, with_all=_UNSET
    ) -> list[str]:
        """``with_all``: the set's evaluation if the caller already has it (values, or None when it
        failed); evaluated here on first need otherwise."""
        try:
            normalised = {str(k): as_override(v) for k, v in overrides.items()}
        except ValueError:
            normalised = {}
        cache = {} if with_all is self._UNSET else {"with": with_all}
        out: list[str] = []
        for name, raw in overrides.items():
            if only is not None and name not in only:
                continue  # filtered BEFORE anything is evaluated
            try:
                value = as_override(raw)
            except ValueError:
                continue
            if isinstance(value, str):
                # A formula override that is the parameter's own formula (spacing aside) is the baseline
                # again. Anything else can only be judged by evaluating it, which the compare does.
                if self.formulas.get(name) and _same_expression(self.formulas[name] or "", value):
                    out.append(name)
            elif name in self.baseline and abs(value - self.baseline[name]) <= abs(self.baseline[name]) * 1e-9:
                if not self.formulas.get(name):
                    out.append(name)  # a plain value set to itself
                    continue
                if "with" not in cache:
                    cache["with"] = self.evaluate(normalised)[0]
                if _changes_nothing(self.database, normalised, name, cache["with"], self.namespace):
                    out.append(name)
        return sorted(out)

    def evaluation_error(self, overrides: dict, values_and_error=None) -> str | None:
        if not self.database or not overrides:
            return None
        values, why = values_and_error if values_and_error is not None else self.evaluate(overrides)
        if why is not None and self.baseline_broken():
            return None  # the model's own finding, not this scenario's
        return why


def judge_all(database: str, entries: list[dict]) -> dict[str, dict]:
    """``ineffective`` and ``evaluation_error`` for every scenario of a list, from ONE read of the model:
    one namespace, one evaluation per scenario (plus the baseline, only if a scenario fails, and the
    rare "without" evaluation of a number set on a formula parameter at its current value)."""
    out: dict[str, dict] = {}
    with lca._LCA_LOCK:
        judge = _Judge(database)
        for entry in entries:
            overrides = entry.get("overrides") or {}
            evaluated = judge.evaluate(overrides) if (database and overrides) else ({}, None)
            out[entry["id"]] = {
                "ineffective": judge.ineffective(overrides, with_all=evaluated[0]),
                "evaluation_error": judge.evaluation_error(overrides, evaluated),
            }
    return out


def _changes_nothing(
    database: str, overrides: dict, name: str, with_it: dict[str, float] | None, namespace
) -> bool | None:
    """Whether dropping the override on formula parameter ``name`` leaves every parameter as it is —
    True, False, or None when the set cannot be evaluated (then nothing is claimed either way, and
    :func:`evaluation_error` is what says why).

    A number on a formula parameter also takes its formula out, so "the same value as now" does not
    prove it changes nothing: with ``a = 1, b = "2*a"``, ``{a: 3, b: 2}`` holds ``b`` at 2 while ``a``
    moves, and without the override ``b`` would follow ``a`` to 6."""
    if with_it is None:
        return None
    try:
        without = params.evaluate_overrides(database, {k: v for k, v in overrides.items() if k != name}, namespace)
    except PrismError:
        return None
    return all(abs(with_it[k] - without.get(k, with_it[k])) <= abs(with_it[k]) * 1e-9 + 1e-12 for k in with_it)


def _evaluation_error(database: str, overrides: dict, namespace=None) -> str | None:
    try:
        normalised = {str(k): as_override(v) for k, v in (overrides or {}).items()}
        params.evaluate_overrides(database, normalised, namespace)
    except ValueError as exc:
        return str(exc)
    except PrismError as exc:
        return exc.message
    return None


def evaluation_error(database: str, overrides: dict) -> str | None:
    """Why the model cannot be evaluated under this scenario's overrides, or None.

    The write paths refuse such a set; this is for the reading side — a scenario stored before, or one
    the baseline has since moved under — so the scenario list can say it instead of listing nothing.
    None too when the model cannot be evaluated even without the overrides: that is the model's
    finding, not the scenario's.

    Known limit, documented rather than fixed: "the baseline is already broken"
    also hides an error the scenario introduces on top — the write path accepts a numeric set then too.
    """
    if not database or not overrides:
        return None
    with lca._LCA_LOCK:
        return _Judge(database).evaluation_error(overrides)


def unreachable_overrides(database: str, overrides: dict, dead: dict | None = None) -> dict[str, list[str]]:
    """Overrides on a parameter that can reach no exchange formula — a different value, no effect.

    The sibling of :func:`ineffective_overrides`, and the one it cannot see. That check compares the
    override against the baseline *value*; these overrides pass it easily, because the value really
    is different. What is missing is the other end: nothing in the model consumes the parameter, so
    the number moves and the result does not.

    A balcony-solar model held three of them. `tilt_factor` 1.16 against a baseline
    of 1.0, `pvgis_yield_ref_kwh_per_kwp` 450 and 650 against 550 — all three plainly different, all
    three returning Δ 0, because the chain they feed ends at `lifetime_energy_kwh` and no exchange
    formula names it. That is the same three-identical-bars symptom the value check was built
    against, reached by the other road (V236).
    """
    from prism_lca.engine import params

    if dead is None:  # ``dead``: the model's answer, for a caller judging many scenarios at once
        dead = params.unreachable_parameters(database)
    return {name: dead[name] for name in (overrides or {}) if name in dead}


def refuse_ineffective(
    database: str, overrides: dict, *, scenario: str = "", only: set[str] | None = None
) -> None:
    """Refuse a scenario that would change nothing. Raises, or returns quietly.

    Only the value half refuses: an override set to the value the model already has is never
    anything but the baseline renamed. The reachability half (:func:`unreachable_overrides`) is
    reported instead — see the comment below.
    """
    from prism_lca.errors import PrismError

    where = f" in scenario '{scenario}'" if scenario else ""
    # ``only``: judge just these names, but within the whole set — a single override set into an
    # existing scenario (set_value) is ineffective or not depending on the others beside it.
    same = ineffective_overrides(database, overrides, only=only)
    if same:
        names = ", ".join(f"'{n}'" for n in same)
        raise PrismError(
            f"{names}{where} would be set to the value the model already has, so this variant would "
            "return the baseline result.",
            code="scenario_override_is_baseline",
            hint="Give the override a different value, or drop it and let the parameter speak for itself.",
        )

    # The reachability half is deliberately NOT refused here. "Same value as the baseline" is never
    # a legitimate intermediate state, but "not wired up yet" is: parameters, scenarios and formulas
    # get written in whatever order the work runs, and a scenario created before the formula it will
    # move is ordinary. So it is reported — in the scenario list and in the model health check —
    # where it can be read once the model is finished, rather than refused where it would block a
    # half-built one (V236).


def _restore(database: str, baseline: dict[str, float], scenario_name: str) -> None:
    """Put a scenario's baseline back. The implementation is shared by every parameter window in the
    app — see :func:`param_window.restore_baseline`, which is where the retry and the wording live."""
    param_window.restore_baseline(database, baseline, scenario_name, kind="scenario")


@contextmanager
def applied(store: ScenarioStore, database: str, scenario_id: str | None):
    """Apply a scenario's parameter overrides for the duration of the block, then restore the
    baseline. A no-op when ``scenario_id`` is falsy — callers can wrap unconditionally. Holds the
    LCA lock across apply→body→restore (see compute()); the lock is re-entrant so a body that
    re-acquires it (e.g. lca.contribution) is fine.

    The window is noted on disk before the overrides are written (see engine/param_window.py): this
    block can be held open for minutes by a "write all chapters" run and for a minute by an
    unattended background warm-up, and a process killed inside it would otherwise leave the model
    standing in the scenario with nothing anywhere to say so."""
    if not scenario_id:
        yield
        return
    scenario = store.get(scenario_id, database)
    with lca._LCA_LOCK:
        baseline = params.snapshot_all(database)
        applied_ = plan(scenario, params.snapshot_formulas(database))
        frame = param_window.open_window(
            database, "scenario", scenario["name"], baseline, applied_.numbers, applied_.replaced
        )
        try:
            params.apply_values(database, {**baseline, **applied_.numbers}, applied_.formulas)
            check_applied(database, applied_, scenario["name"])
            yield
        except (lca.CalcCancelled, PrismError):
            raise
        except Exception as exc:  # noqa: BLE001 - clear message instead of an opaque 500
            raise scenario_calc_error(scenario["name"], exc) from exc
        finally:
            param_window.restore_and_close(database, baseline, frame)


def compute(store: ScenarioStore, database: str, code: str, method, scenario_id: str, amount: float = 1.0) -> dict:
    scenario = store.get(scenario_id, database)
    # Hold the LCA lock across apply→compute→restore: the apply mutates this model's parameter
    # state, and concurrent threadpool calcs would otherwise compute against the perturbed values.
    with lca._LCA_LOCK:
        baseline = params.snapshot_all(database)
        applied_ = plan(scenario, params.snapshot_formulas(database))
        frame = param_window.open_window(
            database, "scenario", scenario["name"], baseline, applied_.numbers, applied_.replaced
        )
        try:
            params.apply_values(database, {**baseline, **applied_.numbers}, applied_.formulas)
            check_applied(database, applied_, scenario["name"])
            result = lca.compute_lcia(database, code, method, amount)
            result["scenario"] = scenario["name"]
            if applied_.replaces_formula():
                result["replaces_formula"] = applied_.replaces_formula()
            return result
        except PrismError:
            raise
        except Exception as exc:  # noqa: BLE001 - wrap bw2data/param failures in a clear message
            raise scenario_calc_error(scenario["name"], exc) from exc
        finally:
            param_window.restore_and_close(database, baseline, frame)


#: Below this relative change a scenario row carries a note naming the exchanges its overrides reach.
#: One run created its scenarios on parameters that were still typed as constants on the
#: exchanges, compared them, read −0.1 % / +0.2 %, spent 187 s on a 16-indicator comparison of the same
#: dead variants — and only the model's own doubt caught it. Nothing said a word (V332).
SMALL_DELTA_PCT = 1.0


def _reach_note(usage: dict[str, list[dict]], overrides: dict, unreachable: set[str] | None = None) -> str:
    """Why a variant may have moved so little: which exchange formulas its overrides actually reach.

    ``unreachable`` is `params.unreachable_parameters` — the transitive answer. `usage` only sees a
    parameter named ON an exchange; one that feeds another parameter's formula (a = 100, b = "2*a",
    the exchange uses b) reaches the result just as well, and the note used to call that variant
    unable to move anything while the row beside it showed it moving."""
    names = sorted(overrides or {})
    reached: list[str] = []
    for name in names:
        for use in usage.get(name, []):
            label = f"'{use.get('input')}' in '{use.get('activity_name')}'"
            if label not in reached:
                reached.append(label)
    quoted = ", ".join(f"'{n}'" for n in names)
    if not reached:
        through = [n for n in names if unreachable is not None and n not in unreachable]
        if through:
            listed = ", ".join(f"'{n}'" for n in through)
            return (
                f"Δ below {SMALL_DELTA_PCT:g} % — {listed} reach exchange formulas only through other "
                "parameters (nothing on an exchange names them directly). If those are the intended "
                "levers, the small effect is the finding. If the quantity you meant to vary is still a "
                'typed constant on its exchange, bind it (edit_exchange action="set_formula") and '
                "compare again."
            )
        return (
            f"Δ below {SMALL_DELTA_PCT:g} % — no exchange formula references {quoted}; this variant "
            "cannot move the result. Bind the quantity to the parameter (edit_exchange "
            'action="set_formula") and compare again, or drop the scenario.'
        )
    shown = ", ".join(reached[:4]) + (f" and {len(reached) - 4} more" if len(reached) > 4 else "")
    return (
        f"Δ below {SMALL_DELTA_PCT:g} % — {quoted} reach {len(reached)} exchange formula(s): {shown}. "
        "If the quantity you meant to vary is still a typed constant on its exchange, that constant "
        'does not follow the parameter: bind it (edit_exchange action="set_formula") and compare again. '
        "If those are the intended levers, the small effect is the finding."
    )


def compare_many(
    store: ScenarioStore, database: str, code: str, methods: list, scenario_ids: list[str], amount: float = 1.0
) -> dict:
    """Baseline + scenarios for SEVERAL indicators, each scenario applied once.

    Applying a scenario rewrites the model's parameter state, which invalidates the factorised LCA;
    scoring another indicator on the same state only swaps the characterisation matrix. So the loop
    is subject-major: apply, score every indicator, next subject — (1 + n) factorisations. The tool
    used to fan out one request per indicator, each applying every scenario again: 16 indicators ×
    4 subjects = 64 factorisations, 187 s and 226 s in one run, where `calculate` over the
    same 16 indicators took 3 s (V331).
    """
    methods = [tuple(m) for m in methods]
    # Before the lock and the baseline: an id copied with a slip used to be found out only after the
    # baseline and the scenarios before it had been calculated. Only a check — each scenario is still read
    # right before it is applied, so an edit made while this waited for the lock is the one scored.
    store.require(database, scenario_ids)
    labels = ["(baseline)"]
    overrides_by: list[dict] = [{}]
    replaced_by: list[dict[str, str]] = [{}]
    scores: list[list[float]] = []
    units: list[str] = []
    name = "(baseline)"
    with lca._LCA_LOCK:  # see compute(): applies mutate this model's parameter state
        baseline = params.snapshot_all(database)
        base_formulas = params.snapshot_formulas(database)
        # One frame for the whole comparison, re-stamped per scenario: this loop never passes back
        # through the baseline between two scenarios, so at any instant exactly one set of overrides
        # is in the model — and that is the set the marker has to name.
        frame = param_window.open_window(database, "comparison", "a scenario comparison", baseline, {})
        try:
            params.apply_values(database, baseline, base_formulas)
            first = [lca.compute_lcia(database, code, m, amount) for m in methods]
            units = [r["unit"] for r in first]
            scores.append([r["score"] for r in first])
            for sid in scenario_ids:
                scenario = store.get(sid, database)
                name = scenario["name"]
                applied_ = plan(scenario, base_formulas)
                param_window.set_overrides(frame, applied_.numbers, applied_.replaced)
                params.apply_values(database, {**baseline, **applied_.numbers}, applied_.formulas)
                param_window.settle(frame)
                check_applied(database, applied_, name)
                scores.append([lca.compute_lcia(database, code, m, amount)["score"] for m in methods])
                labels.append(name)
                overrides_by.append({**applied_.numbers, **applied_.formula_overrides})
                replaced_by.append(applied_.replaces_formula())
        except PrismError:
            raise
        except Exception as exc:  # noqa: BLE001 — wrap bw2data/param failures in a clear message
            raise scenario_calc_error(name, exc) from exc
        finally:
            param_window.restore_and_close(database, baseline, frame)
    # The difference to the baseline is the whole point of the comparison, and it was left to the
    # caller to divide — so the "-21 %" in a written answer was arithmetic done by hand, outside
    # anything that can be checked. Absolute and relative, with the relative one omitted rather than
    # faked when the baseline is zero.
    usage: dict[str, list[dict]] | None = None
    unreachable: set[str] | None = None
    out: list[dict] = []
    for j, m in enumerate(methods):
        results = [{"scenario": labels[i], "score": scores[i][j]} for i in range(len(labels))]
        base_score = results[0]["score"]
        for i, row in enumerate(results[1:], start=1):
            row["delta"] = row["score"] - base_score
            # An override on a parameter that has a formula of its own replaces that formula for the
            # variant — named, so whoever reads "b = 6" sees that "b = 2*a" is what it stood in for.
            if replaced_by[i]:
                row["replaces_formula"] = replaced_by[i]
            if base_score:
                row["delta_pct"] = 100.0 * (row["score"] - base_score) / abs(base_score)
                if abs(row["delta_pct"]) < SMALL_DELTA_PCT:
                    if usage is None:
                        usage = params.formula_usage(database)
                        unreachable = set(params.unreachable_parameters(database))
                    row["note"] = _reach_note(usage, overrides_by[i], unreachable)
        out.append(
            {"database": database, "code": code, "method": list(m), "amount": amount,
             "unit": units[j], "baseline_score": base_score, "results": results}
        )
    return {"database": database, "code": code, "amount": amount, "results": out}


def compare(store: ScenarioStore, database: str, code: str, method, scenario_ids: list[str], amount: float = 1.0) -> dict:
    """Baseline + scenarios for ONE indicator — `compare_many` with a single method."""
    return compare_many(store, database, code, [method], scenario_ids, amount)["results"][0]


# ------------------------------------------------------------------------- #
# Excel import / template  (rows = parameters, columns = scenarios)
# ------------------------------------------------------------------------- #
def template_workbook(database: str) -> bytes:
    """A ready-to-edit example workbook: parameters down the first column, scenario names across
    the first row. ``Scenario A`` is prefilled with the current baseline amounts; ``Scenario B`` is
    left blank to show that an empty cell keeps the baseline (only filled cells override)."""
    rows = params.list_params(database)
    wb = Workbook()
    ws = wb.active
    ws.title = "scenarios"
    ws.append(["parameter", "Scenario A", "Scenario B"])
    for p in rows:
        base = p.get("amount")
        ws.append([p["name"], base if base is not None else None, None])
    if not rows:  # no parameters yet — still show the shape with a placeholder row
        ws.append(["example_parameter", 1, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_scenarios_workbook(content: bytes) -> list[dict]:
    """Parse a scenario workbook into ordered columns ``[{"name", "overrides"}]``.

    First row (from the 2nd cell on) = scenario names; first column (from the 2nd row on) =
    parameter names. A blank cell means "don't override" (skip it); a numeric cell — including 0 —
    overrides; a text cell is a formula override (checked by the import against the model, like every
    other formula override)."""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    grid = list(ws.iter_rows(values_only=True))
    if not grid:
        raise PrismError("The workbook is empty.", code="empty_import",
                         hint="Use the template button to get a correctly shaped file.")

    header = grid[0] or ()
    scen_cols: list[tuple[int, str]] = []
    for j, cell in enumerate(header):
        if j == 0:
            continue  # top-left corner is the parameter-column label
        name = str(cell).strip() if cell is not None else ""
        if name:
            scen_cols.append((j, name))
    if not scen_cols:
        raise PrismError("No scenario columns found in the first row.", code="no_scenarios",
                         hint="Put scenario names in the first row, from the second column onward.")

    columns = {j: {"name": name, "overrides": {}} for j, name in scen_cols}
    for row in grid[1:]:
        if row is None or all(c is None for c in row):
            continue
        pname = str(row[0]).strip() if row and row[0] is not None else ""
        if not pname:
            continue
        for j, sname in scen_cols:
            cell = row[j] if j < len(row) else None
            if cell is None or (isinstance(cell, str) and not cell.strip()):
                continue  # empty → keep baseline / existing value
            try:
                if isinstance(cell, str):
                    try:
                        value: float | str = float(cell.strip().replace(",", "."))
                    except ValueError:
                        value = as_override(cell)  # a formula; the import evaluates it
                else:
                    value = float(cell)
            except (TypeError, ValueError) as exc:
                raise PrismError(
                    f"Cell for parameter '{pname}' in scenario '{sname}' is not a number or a formula: {cell!r}.",
                    code="invalid_cell",
                    hint="Cells must be numbers or formulas over the model's parameters; leave a cell "
                    "empty to keep the baseline value.") from exc
            columns[j]["overrides"][pname] = value
    return [columns[j] for j, _ in scen_cols]
